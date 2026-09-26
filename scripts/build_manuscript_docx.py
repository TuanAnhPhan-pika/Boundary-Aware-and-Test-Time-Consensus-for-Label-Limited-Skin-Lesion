"""Build the anonymous journal-style DOCX from the English Markdown manuscript."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper" / "paper_english.md"
OUTPUT = ROOT / "paper" / "submission_manuscript.docx"
NAVY = "17365D"
PALE = "EAF1F8"
GRID = "D9D9D9"


def set_font(run, name="Times New Roman", size=10.5, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_inline(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2]); set_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1]); set_font(run, name="Consolas", size=9)
        else:
            run = paragraph.add_run(part); set_font(run)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar"); tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}"); margins.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def borders(table):
    tbl_pr = table._tbl.tblPr
    edges = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single"); tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), GRID); edges.append(tag)
    tbl_pr.append(edges)


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    borders(table)
    for j, value in enumerate(rows[0]):
        cell = table.rows[0].cells[j]; shade(cell, NAVY); cell_margins(cell)
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(value.strip()); set_font(run, size=8.2, bold=True); run.font.color.rgb = RGBColor(255,255,255)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for i, source_row in enumerate(rows[1:], start=1):
        cells = table.add_row().cells
        for j, value in enumerate(source_row):
            cell = cells[j]; cell_margins(cell)
            if i % 2 == 0: shade(cell, PALE)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            cleaned = value.strip().replace("**", "")
            run = p.add_run(cleaned); set_font(run, size=8.1)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar"); fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar"); fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2]); set_font(run, size=9)


def style_document(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.75); sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.82); sec.right_margin = Inches(0.82)
    sec.header_distance = Inches(0.3); sec.footer_distance = Inches(0.3)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(5)
    for style_name, size, before, after in (("Title", 17, 0, 12), ("Heading 1", 13, 12, 5), ("Heading 2", 11, 9, 3)):
        style = doc.styles[style_name]
        style.font.name = "Arial"; style.font.size = Pt(size); style.font.bold = True
        style.font.color.rgb = RGBColor(0,0,0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style_p_pr = style._element.get_or_add_pPr()
        style_border = style_p_pr.find(qn("w:pBdr"))
        if style_border is not None:
            style_p_pr.remove(style_border)
    doc.styles["Caption"].font.name = "Times New Roman"
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.italic = True
    doc.styles["Caption"].font.color.rgb = RGBColor(0,0,0)


def clean_heading(text):
    return re.sub(r"^\d+(?:\.\d+)*\.?\s*", "", text).strip().rstrip(".?:")


def build():
    doc = Document(); style_document(doc)
    header = doc.sections[0].header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("ISIC 2016 segmentation study"); set_font(run, name="Arial", size=8)
    add_page_number(doc.sections[0].footer.paragraphs[0])

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    figure_number = 0
    table_number = 0
    in_references = False
    table_captions = [
        "Primary endpoint across label budgets",
        "Paired Dice differences relative to the compact U-Net",
        "Secondary endpoints at the full label budget",
    ]
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, line[2:])
            p_pr = p._p.get_or_add_pPr()
            p_bdr = p_pr.find(qn("w:pBdr"))
            if p_bdr is not None:
                p_pr.remove(p_bdr)
            p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p2.add_run("Anonymous manuscript for peer review"); set_font(run, size=10, italic=True)
            i += 1; continue
        if line.startswith("## "):
            heading = clean_heading(line[3:])
            in_references = heading.lower() == "references"
            doc.add_paragraph(heading, style="Heading 1"); i += 1; continue
        if line.startswith("### "):
            doc.add_paragraph(clean_heading(line[4:]), style="Heading 2"); i += 1; continue
        if line.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            if match:
                alt, rel = match.groups(); image = (SOURCE.parent / rel).resolve()
                doc.add_page_break()
                figure_number += 1
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                shape = p.add_run().add_picture(str(image), width=Inches(6.35))
                shape._inline.docPr.set("descr", alt)
                cap = doc.add_paragraph(style="Caption"); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.add_run(f"Figure {figure_number}. {alt}.")
                if figure_number == 3:
                    doc.add_page_break()
            i += 1; continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip()); i += 1
            rows = [[c.strip() for c in row.strip("|").split("|")] for row in block]
            if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
                rows.pop(1)
            table_number += 1
            caption = table_captions[table_number - 1] if table_number <= len(table_captions) else "Results table"
            cp = doc.add_paragraph(style="Caption"); cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            cp.add_run(f"Table {table_number}. {caption}.")
            add_table(doc, rows); continue
        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s*", "", line))
            if in_references:
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Inches(0.18)
                p.paragraph_format.first_line_indent = Inches(-0.18)
                for run in p.runs:
                    set_font(run, size=8.6)
            i += 1; continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); add_inline(p, line[2:]); i += 1; continue
        parts = [line]; i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||!\[|\d+\.\s|-\s)", lines[i].strip()):
            parts.append(lines[i].strip()); i += 1
        p = doc.add_paragraph(); add_inline(p, " ".join(parts))
        if parts[0].startswith("**Keywords:"):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    props = doc.core_properties
    props.title = "Boundary Loss and Test Time Consensus in Label Limited ISIC 2016 Skin Lesion Segmentation"
    props.subject = "Anonymous research manuscript"
    props.author = "Anonymous"
    props.keywords = "ISIC 2016, skin lesion segmentation, U-Net, test-time augmentation, boundary loss"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
