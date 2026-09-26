"""Build the Vietnamese research-explanation DOCX from Markdown."""

from __future__ import annotations

import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches

from build_manuscript_docx import (
    ROOT,
    add_inline,
    add_page_number,
    add_table,
    clean_heading,
    set_font,
    style_document,
)


SOURCE = ROOT / "paper" / "thuyet_minh_nghien_cuu_tieng_viet.md"
OUTPUT = ROOT / "paper" / "thuyet_minh_nghien_cuu_tieng_viet.docx"


def build() -> None:
    doc = Document()
    style_document(doc)

    header = doc.sections[0].header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("Thuyết minh nghiên cứu ISIC 2016"), name="Arial", size=8)
    add_page_number(doc.sections[0].footer.paragraphs[0])

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    table_captions = [
        "Thông tin chung của đề tài",
        "Bảy cấu hình được đánh giá",
        "Dice trung bình theo lượng dữ liệu huấn luyện",
        "Sản phẩm của đề tài",
    ]
    figure_captions = [
        "Dice của các phương pháp theo lượng dữ liệu huấn luyện",
        "Chênh lệch Dice so với U-Net gọn nhẹ",
    ]
    table_number = 0
    figure_number = 0
    i = 0

    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        if line.startswith("# "):
            title = doc.add_paragraph(style="Title")
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(title, line[2:])
            subtitle = doc.add_paragraph()
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(
                subtitle.add_run("Bản thuyết minh kết quả nghiên cứu"),
                size=10,
                italic=True,
            )
            i += 1
            continue

        if line.startswith("## "):
            doc.add_paragraph(clean_heading(line[3:]), style="Heading 1")
            i += 1
            continue

        if line.startswith("### "):
            doc.add_paragraph(clean_heading(line[4:]), style="Heading 2")
            i += 1
            continue

        if line.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            if match:
                alt, rel = match.groups()
                image_path = (SOURCE.parent / rel).resolve()
                doc.add_page_break()
                figure_number += 1
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                shape = paragraph.add_run().add_picture(str(image_path), width=Inches(6.25))
                shape._inline.docPr.set("descr", alt)
                caption = doc.add_paragraph(style="Caption")
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption_text = figure_captions[figure_number - 1] if figure_number <= len(figure_captions) else alt
                caption.add_run(f"Hình {figure_number}. {caption_text}.")
            i += 1
            continue

        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = [[cell.strip() for cell in row.strip("|").split("|")] for row in block]
            if len(rows) > 1 and all(
                re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in rows[1]
            ):
                rows.pop(1)
            table_number += 1
            caption_text = (
                table_captions[table_number - 1]
                if table_number <= len(table_captions)
                else "Bảng kết quả"
            )
            caption = doc.add_paragraph(style="Caption")
            caption.add_run(f"Bảng {table_number}. {caption_text}.")
            add_table(doc, rows)
            continue

        if re.match(r"^\d+\.\s", line):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.first_line_indent = Inches(-0.22)
            add_inline(paragraph, line)
            i += 1
            continue

        if line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
            i += 1
            continue

        parts = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^(#|\||!\[|\d+\.\s|-\s)", lines[i].strip()
        ):
            parts.append(lines[i].strip())
            i += 1
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(paragraph, " ".join(parts))

    properties = doc.core_properties
    properties.title = "Thuyết minh đề tài phân đoạn tổn thương da khi dữ liệu được vẽ mặt nạ còn hạn chế"
    properties.subject = "Bản thuyết minh kết quả nghiên cứu xử lý ảnh y khoa"
    properties.author = "Nhóm nghiên cứu"
    properties.keywords = "ISIC 2016, phân đoạn tổn thương da, U-Net, TTA, loss vùng biên"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
